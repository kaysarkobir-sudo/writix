"""
Bulk Book Import Views
Add these functions to your student_management/views.py

Features:
- Download Excel/CSV template with correct format
- Upload and validate bulk book data
- Preview before importing
- Error reporting with line numbers
- Support for Excel (.xlsx) and CSV formats
"""

import csv
import io
from datetime import date
from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse
from decimal import Decimal, InvalidOperation

# For Excel support
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False


@login_required
def download_book_import_template(request):
    """
    Download Excel/CSV template for bulk book import
    URL: /library/books/import/template/
    """

    # Get format from query parameter (default to Excel if available)
    file_format = request.GET.get('format', 'excel' if EXCEL_AVAILABLE else 'csv')

    # Define template headers and sample data
    headers = [
        'name',           # Book title (required)
        'author',         # Author name (required)
        'publisher',      # Publisher name
        'isbn',           # ISBN number
        'edition',        # Edition (e.g., 1st, 2nd)
        'publication_year',  # Year (e.g., 2024)
        'category',       # Category (fiction, science, etc.)
        'subject_code',   # Subject code (e.g., CS101)
        'language',       # Language (e.g., English)
        'price',          # Price in BDT (required)
        'quantity',       # Total quantity (required)
        'available_quantity',  # Available (same as quantity initially)
        'rack_no',        # Rack number (e.g., A1, B2)
        'shelf_location', # Shelf location description
        'pages',          # Number of pages
        'description',    # Brief description
        'barcode',        # Barcode number
    ]

    # Sample data row
    sample_data = [
        'Introduction to Python Programming',
        'John Smith',
        'Tech Publishers',
        '978-0-123456-78-9',
        '3rd Edition',
        '2024',
        'science',
        'CS101',
        'English',
        '850.00',
        '10',
        '10',
        'A1',
        'Computer Science Section - Top Shelf',
        '450',
        'Comprehensive guide to Python programming for beginners',
        'BK2024001',
    ]

    instructions = [
        'INSTRUCTIONS:',
        '1. Fill in the data rows below',
        '2. Required fields: name, author, price, quantity',
        '3. Category options: fiction, non_fiction, science, mathematics, history, geography, literature, biography, reference, other',
        '4. Leave available_quantity same as quantity for new books',
        '5. Date format: YYYY (e.g., 2024)',
        '6. Price: numbers only (e.g., 850.00)',
        '7. Delete this instruction row before uploading',
        '',
    ]

    if file_format == 'excel' and EXCEL_AVAILABLE:
        # Create Excel file
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Book Import Template'

        # Add instructions
        for i, instruction in enumerate(instructions, 1):
            cell = ws.cell(row=i, column=1, value=instruction)
            cell.font = Font(bold=True, color='FF0000')

        # Add headers
        header_row = len(instructions) + 1
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col, value=header)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')

        # Add sample data
        for col, value in enumerate(sample_data, 1):
            cell = ws.cell(row=header_row + 1, column=col, value=value)
            cell.fill = PatternFill(start_color='E7E6E6', end_color='E7E6E6', fill_type='solid')

        # Adjust column widths
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width

        # Save to response
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="book_import_template_{date.today()}.xlsx"'
        wb.save(response)

    else:
        # Create CSV file
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="book_import_template_{date.today()}.csv"'

        writer = csv.writer(response)

        # Add instructions as comments
        for instruction in instructions:
            writer.writerow([instruction])

        # Add headers
        writer.writerow(headers)

        # Add sample data
        writer.writerow(sample_data)

    return response


@login_required
def bulk_book_import(request):
    """
    Upload and import books from Excel/CSV file
    URL: /library/books/import/
    """

    if request.method == 'POST':
        # Check if file was uploaded
        if 'file' not in request.FILES:
            messages.error(request, 'Please select a file to upload.')
            return redirect('bulk_book_import')

        uploaded_file = request.FILES['file']

        # Check file extension
        file_name = uploaded_file.name.lower()
        if not (file_name.endswith('.xlsx') or file_name.endswith('.csv')):
            messages.error(request, 'Please upload an Excel (.xlsx) or CSV (.csv) file.')
            return redirect('bulk_book_import')

        # Parse file
        try:
            if file_name.endswith('.xlsx'):
                books_data, errors = parse_excel_file(uploaded_file)
            else:
                books_data, errors = parse_csv_file(uploaded_file)

            # If there are parsing errors, show them
            if errors:
                context = {
                    'errors': errors,
                    'error_count': len(errors),
                }
                return render(request, 'student_management/library/bulk_import_errors.html', context)

            # Preview mode or final import
            action = request.POST.get('action', 'preview')

            if action == 'preview':
                # Show preview
                context = {
                    'books_data': books_data,
                    'total_books': len(books_data),
                }
                return render(request, 'student_management/library/bulk_import_preview.html', context)

            elif action == 'import':
                # Import books
                success_count = 0
                error_count = 0
                errors = []

                for idx, book_data in enumerate(books_data, 1):
                    try:
                        # Create book
                        Book.objects.create(**book_data)
                        success_count += 1
                    except Exception as e:
                        error_count += 1
                        errors.append(f"Row {idx}: {str(e)}")

                if success_count > 0:
                    messages.success(request, f'Successfully imported {success_count} books!')

                if error_count > 0:
                    messages.warning(request, f'{error_count} books failed to import. Check errors below.')
                    context = {
                        'errors': errors,
                        'error_count': error_count,
                        'success_count': success_count,
                    }
                    return render(request, 'student_management/library/bulk_import_errors.html', context)

                return redirect('book_list')

        except Exception as e:
            messages.error(request, f'Error processing file: {str(e)}')
            return redirect('bulk_book_import')

    # GET request - show upload form
    context = {
        'excel_available': EXCEL_AVAILABLE,
    }
    return render(request, 'student_management/library/bulk_import_form.html', context)


def parse_excel_file(uploaded_file):
    """Parse Excel file and return book data"""

    if not EXCEL_AVAILABLE:
        raise Exception('openpyxl is not installed. Please install it: pip install openpyxl')

    books_data = []
    errors = []

    try:
        wb = openpyxl.load_workbook(uploaded_file)
        ws = wb.active

        # Find header row (skip instruction rows)
        header_row = None
        for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=20), 1):
            if row[0].value and str(row[0].value).lower() == 'name':
                header_row = row_idx
                break

        if not header_row:
            errors.append('Header row not found. Make sure first column header is "name"')
            return books_data, errors

        # Get headers
        headers = [cell.value for cell in ws[header_row] if cell.value]

        # Process data rows
        for row_idx, row in enumerate(ws.iter_rows(min_row=header_row + 1), header_row + 1):
            # Skip empty rows
            if not row[0].value:
                continue

            # Skip instruction/sample rows
            if str(row[0].value).startswith('INSTRUCTIONS') or str(row[0].value).startswith('Introduction to Python'):
                continue

            # Create book data dictionary
            book_data = {}
            row_errors = []

            for col_idx, header in enumerate(headers):
                cell_value = row[col_idx].value if col_idx < len(row) else None

                # Clean and validate
                if header == 'name':
                    if not cell_value:
                        row_errors.append(f'Row {row_idx}: Book name is required')
                    else:
                        book_data['name'] = str(cell_value).strip()

                elif header == 'author':
                    if not cell_value:
                        row_errors.append(f'Row {row_idx}: Author is required')
                    else:
                        book_data['author'] = str(cell_value).strip()

                elif header == 'price':
                    if not cell_value:
                        row_errors.append(f'Row {row_idx}: Price is required')
                    else:
                        try:
                            book_data['price'] = Decimal(str(cell_value))
                        except (InvalidOperation, ValueError):
                            row_errors.append(f'Row {row_idx}: Invalid price format')

                elif header == 'quantity':
                    if not cell_value:
                        row_errors.append(f'Row {row_idx}: Quantity is required')
                    else:
                        try:
                            book_data['quantity'] = int(cell_value)
                        except (ValueError, TypeError):
                            row_errors.append(f'Row {row_idx}: Invalid quantity format')

                elif header == 'available_quantity':
                    if cell_value:
                        try:
                            book_data['available_quantity'] = int(cell_value)
                        except (ValueError, TypeError):
                            row_errors.append(f'Row {row_idx}: Invalid available_quantity format')
                    else:
                        book_data['available_quantity'] = book_data.get('quantity', 1)

                elif header == 'publication_year':
                    if cell_value:
                        try:
                            book_data['publication_year'] = int(cell_value)
                        except (ValueError, TypeError):
                            row_errors.append(f'Row {row_idx}: Invalid publication year')

                elif header == 'pages':
                    if cell_value:
                        try:
                            book_data['pages'] = int(cell_value)
                        except (ValueError, TypeError):
                            row_errors.append(f'Row {row_idx}: Invalid pages format')

                elif header in ['publisher', 'isbn', 'edition', 'category', 'subject_code',
                               'language', 'rack_no', 'shelf_location', 'description', 'barcode']:
                    if cell_value:
                        book_data[header] = str(cell_value).strip()

            if row_errors:
                errors.extend(row_errors)
            elif book_data:
                books_data.append(book_data)

    except Exception as e:
        errors.append(f'Error parsing Excel file: {str(e)}')

    return books_data, errors


def parse_csv_file(uploaded_file):
    """Parse CSV file and return book data"""

    books_data = []
    errors = []

    try:
        # Decode file
        file_data = uploaded_file.read().decode('utf-8')
        csv_data = csv.reader(io.StringIO(file_data))

        # Find header row
        header_row = None
        headers = None

        for row_idx, row in enumerate(csv_data, 1):
            if row and row[0].lower() == 'name':
                headers = [col.strip() for col in row]
                header_row = row_idx
                break

        if not headers:
            errors.append('Header row not found. Make sure first column header is "name"')
            return books_data, errors

        # Reset file pointer
        uploaded_file.seek(0)
        file_data = uploaded_file.read().decode('utf-8')
        csv_data = csv.DictReader(io.StringIO(file_data), fieldnames=headers)

        # Skip to data rows
        for _ in range(header_row):
            next(csv_data)

        # Process data rows
        for row_idx, row in enumerate(csv_data, header_row + 1):
            # Skip empty or instruction rows
            if not row.get('name') or row['name'].startswith('INSTRUCTIONS'):
                continue

            book_data = {}
            row_errors = []

            # Required fields
            if row.get('name'):
                book_data['name'] = row['name'].strip()
            else:
                row_errors.append(f'Row {row_idx}: Book name is required')

            if row.get('author'):
                book_data['author'] = row['author'].strip()
            else:
                row_errors.append(f'Row {row_idx}: Author is required')

            if row.get('price'):
                try:
                    book_data['price'] = Decimal(row['price'])
                except (InvalidOperation, ValueError):
                    row_errors.append(f'Row {row_idx}: Invalid price format')
            else:
                row_errors.append(f'Row {row_idx}: Price is required')

            if row.get('quantity'):
                try:
                    book_data['quantity'] = int(row['quantity'])
                except (ValueError, TypeError):
                    row_errors.append(f'Row {row_idx}: Invalid quantity format')
            else:
                row_errors.append(f'Row {row_idx}: Quantity is required')

            # Optional fields
            if row.get('available_quantity'):
                try:
                    book_data['available_quantity'] = int(row['available_quantity'])
                except (ValueError, TypeError):
                    pass
            else:
                book_data['available_quantity'] = book_data.get('quantity', 1)

            if row.get('publication_year'):
                try:
                    book_data['publication_year'] = int(row['publication_year'])
                except (ValueError, TypeError):
                    pass

            if row.get('pages'):
                try:
                    book_data['pages'] = int(row['pages'])
                except (ValueError, TypeError):
                    pass

            # Text fields
            for field in ['publisher', 'isbn', 'edition', 'category', 'subject_code',
                         'language', 'rack_no', 'shelf_location', 'description', 'barcode']:
                if row.get(field):
                    book_data[field] = row[field].strip()

            if row_errors:
                errors.extend(row_errors)
            elif book_data:
                books_data.append(book_data)

    except Exception as e:
        errors.append(f'Error parsing CSV file: {str(e)}')

    return books_data, errors
